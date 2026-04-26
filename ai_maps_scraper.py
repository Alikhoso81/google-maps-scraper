# ============================================================
#   Google Maps Lead Scraper — FINAL VERSION
#   No API key needed. Uses Playwright to browse like a human.
# ============================================================
#
# SETUP (run these once in terminal before first use):
#   pip install playwright openpyxl
#   playwright install chromium
#
# USAGE EXAMPLE:
# python ai_maps_scraper.py --keyword "dentists" --location "Los Angeles California" --max 120 --output "D:\all_leads.xlsx"


import asyncio
import argparse
import re
import os
from datetime import datetime
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCROLL_PAUSE = 1.5

async def scrape_google_maps(keyword, location, max_results, output):
    query = f"{keyword} in {location}"
    print(f"\n🔍  Keyword  : {keyword}")
    print(f"📍  Location : {location}")
    print(f"📦  Target   : {max_results} results")
    print(f"💾  Output   : {output}\n")

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,
            slow_mo=30,
            args=["--no-sandbox","--disable-blink-features=AutomationControlled","--start-maximized"]
        )
        context = await browser.new_context(
            viewport={"width": 1366, "height": 900},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
            locale="en-US",
        )
        page = await context.new_page()
        url = f"https://www.google.com/maps/search/{query.replace(' ', '+')}/"
        print(f"🌐  Opening Google Maps...")

        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=60000)
        except PlaywrightTimeout:
            try:
                await page.goto(url, wait_until="commit", timeout=90000)
            except:
                print("❌  Cannot reach Google Maps.")
                await browser.close()
                return []

        await asyncio.sleep(3)

        for btn in ["Accept all", "Accept", "Agree", "I agree"]:
            try:
                await page.click(f'button:has-text("{btn}")', timeout=2000)
                await asyncio.sleep(1)
                break
            except:
                pass

        print("⏳  Waiting for results...")
        try:
            await page.wait_for_selector('div[role="feed"]', timeout=20000)
        except:
            try:
                await page.wait_for_selector('a[href*="/maps/place/"]', timeout=15000)
            except:
                print("❌  No results found.")
                await browser.close()
                return []

        print("✅  Results loaded! Collecting listings...\n")

        listing_urls = []
        seen_urls = set()
        stale = 0

        while len(listing_urls) < max_results:
            links = await page.locator('a[href*="/maps/place/"]').all()
            new_found = 0
            for link in links:
                href = await link.get_attribute("href") or ""
                if href and href not in seen_urls:
                    seen_urls.add(href)
                    listing_urls.append(href)
                    new_found += 1
            print(f"   📍 {len(listing_urls)} listings collected...", end="\r")
            if new_found == 0:
                stale += 1
                if stale >= 4:
                    print(f"\n   ↳ End of results ({len(listing_urls)} total)")
                    break
            else:
                stale = 0
            try:
                await page.evaluate("const feed = document.querySelector('div[role=\"feed\"]'); if (feed) feed.scrollTop += 2500;")
            except:
                pass
            await asyncio.sleep(SCROLL_PAUSE)

        listing_urls = listing_urls[:max_results]
        print(f"\n📋  Visiting {len(listing_urls)} listings...\n")

        leads = []

        for i, href in enumerate(listing_urls):
            try:
                await page.goto(href, wait_until="domcontentloaded", timeout=30000)
                await asyncio.sleep(1.5)
            except:
                print(f"   [{i+1:>3}] ⚠️  Skipped")
                continue

            lead = {
                "Business Name": "",
                "Phone": "",
                "Website": "",
                "Google Maps URL": page.url,
                "Address": "",
                "City": "",
                "State": "",
                "Rating": "",
                "Total Reviews": "",
                "Business Type": keyword.title(),
                "Hours": "",
                "Owner Name": "",
                "Owner Email": "",
                "Scraped Date": datetime.today().strftime("%Y-%m-%d"),
                "Notes": "",
            }

            for sel in ['h1.DUwDvf', 'h1[class*="DUwDvf"]', 'h1']:
                try:
                    name = await page.locator(sel).first.inner_text(timeout=2000)
                    if name.strip():
                        lead["Business Name"] = name.strip()
                        break
                except:
                    pass

            try:
                phone = await page.locator('[data-item-id*="phone"],[data-tooltip*="phone"]').first.inner_text(timeout=2000)
                if re.search(r'\d{3,}', phone.strip()):
                    lead["Phone"] = phone.strip()
            except:
                pass

            if not lead["Phone"]:
                try:
                    els = await page.locator('button[aria-label*="phone"], a[aria-label*="phone"]').all()
                    for el in els:
                        txt = await el.inner_text(timeout=1000)
                        if re.search(r'\d{3,}', txt):
                            lead["Phone"] = txt.strip()
                            break
                except:
                    pass

            if not lead["Phone"]:
                try:
                    content = await page.content()
                    phones = re.findall(r'(\+1[\s\-]?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{4}|\(\d{3}\)\s?\d{3}[\-\s]\d{4}|\d{3}[\-\.]\d{3}[\-\.]\d{4})', content)
                    if phones:
                        lead["Phone"] = phones[0].replace('&#160;', ' ').replace('&nbsp;', ' ').strip()
                except:
                    pass

            if not lead["Phone"]:
                try:
                    tel_links = await page.locator('a[href^="tel:"]').all()
                    for tl in tel_links:
                        href_tel = await tl.get_attribute("href") or ""
                        number = href_tel.replace("tel:", "").strip()
                        if number:
                            lead["Phone"] = number
                            break
                except:
                    pass

            try:
                web = await page.locator('a[data-item-id="authority"]').first.get_attribute("href", timeout=2000)
                lead["Website"] = web or ""
            except:
                try:
                    links2 = await page.locator('a[href^="http"]').all()
                    for lnk in links2:
                        h = await lnk.get_attribute("href") or ""
                        if h.startswith("http") and "google" not in h and "gstatic" not in h:
                            lead["Website"] = h
                            break
                except:
                    pass

            for sel in ['button[data-item-id*="address"] div.fontBodyMedium','button[data-tooltip*="address"] div']:
                try:
                    addr = await page.locator(sel).first.inner_text(timeout=2000)
                    addr = addr.strip()
                    if addr and len(addr) > 5:
                        lead["Address"] = addr
                        parts = [p.strip() for p in addr.split(",")]
                        if len(parts) >= 3:
                            lead["City"] = parts[-3]
                            lead["State"] = re.sub(r'\d+', '', parts[-2]).strip()
                        elif len(parts) == 2:
                            lead["City"] = parts[0]
                        break
                except:
                    pass

            for sel in ['div.fontDisplayLarge', 'span.MW4etd']:
                try:
                    r = await page.locator(sel).first.inner_text(timeout=1000)
                    if re.match(r'^\d[\.,]\d$', r.strip()):
                        lead["Rating"] = r.strip()
                        break
                except:
                    pass

            for sel in ['button[jsaction*="reviewChart"] span', 'span[aria-label*="review"]']:
                try:
                    rv = await page.locator(sel).first.inner_text(timeout=1000)
                    clean = re.sub(r'[^\d]', '', rv)
                    if clean:
                        lead["Total Reviews"] = clean
                        break
                except:
                    pass

            for sel in ['div[class*="o0Svhf"]', 'div[class*="OqCZI"]']:
                try:
                    h2 = await page.locator(sel).first.inner_text(timeout=1000)
                    if h2.strip():
                        lead["Hours"] = h2.strip()[:60]
                        break
                except:
                    pass

            leads.append(lead)
            p_icon = "📞" if lead["Phone"] else "  "
            w_icon = "🌐" if lead["Website"] else "  "
            print(f"   [{len(leads):>3}/{len(listing_urls)}] {p_icon}{w_icon}  {lead['Business Name'][:55]}")

        await browser.close()

    print(f"\n{'─'*55}")
    print(f"📊  Total scraped  : {len(leads)}")
    print(f"📞  With phone     : {sum(1 for r in leads if r['Phone'])}")
    print(f"🌐  With website   : {sum(1 for r in leads if r['Website'])}")
    print(f"{'─'*55}")
    append_excel(leads, output, keyword, location)
    return leads


COLS = ["Business Name","Phone","Website","Google Maps URL","Address","City","State","Business Type","Rating","Total Reviews","Hours","Owner Name","Owner Email","Scraped Date","Notes"]
WIDTHS = [32,18,30,44,34,16,12,18,8,13,22,20,28,14,20]
THIN = Side(style="thin", color="C0D0E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def make_header(ws):
    for ci, col in enumerate(COLS, 1):
        c = ws.cell(1, ci, col)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", start_color="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 22
    for ci, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}1"

def append_excel(leads, path, keyword, location):
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        start_row = ws.max_row + 1
        print(f"\n📂  Appending to existing file ({ws.max_row - 1} rows already)")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "All Leads"
        make_header(ws)
        start_row = 2
        print(f"\n📂  Creating new file: {path}")

    for ri, row in enumerate(leads, start_row):
        bg = "EBF3FB" if ri % 2 == 0 else "FFFFFF"
        for ci, col in enumerate(COLS, 1):
            val = row.get(col, "")
            c = ws.cell(ri, ci, val)
            c.font = Font(name="Calibri", size=9)
            c.border = BORDER
            c.fill = PatternFill("solid", start_color=bg)
            c.alignment = Alignment(horizontal="center" if col in ("Rating","Total Reviews","Scraped Date") else "left", vertical="center")

    wb.save(path)
    print(f"✅  Saved → {path}")
    print(f"📋  Total rows in file now: {ws.max_row - 1}")

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--keyword", required=True)
    parser.add_argument("--location", required=True)
    parser.add_argument("--max", type=int, default=100)
    parser.add_argument("--output", default="maps_leads.xlsx")
    args = parser.parse_args()
    asyncio.run(scrape_google_maps(args.keyword, args.location, args.max, args.output))

if __name__ == "__main__":
    main()